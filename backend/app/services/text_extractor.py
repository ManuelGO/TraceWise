"""Text extraction service for multiple file formats."""

import io
import logging

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """Raised when text extraction fails."""

    pass


def extract_from_pdf(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from PDF file.

    Args:
        file_bytes: Raw PDF file bytes

    Returns:
        Tuple of (extracted_text, page_count)

    Raises:
        ExtractionError: If PDF is corrupted or cannot be read
    """
    try:
        pdf_reader = PdfReader(io.BytesIO(file_bytes))
        page_count = len(pdf_reader.pages)

        text_parts = []
        for page_num, page in enumerate(pdf_reader.pages, start=1):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                else:
                    logger.warning(f"No text extracted from PDF page {page_num}")
            except Exception as e:
                logger.warning(f"Failed to extract text from PDF page {page_num}: {e}")
                text_parts.append(f"[Page {page_num}: extraction failed]")

        full_text = "\n".join(text_parts)
        return full_text, page_count
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from PDF: {e}") from e


def extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file.

    Args:
        file_bytes: Raw DOCX file bytes

    Returns:
        Extracted text (paragraphs + tables)

    Raises:
        ExtractionError: If DOCX is corrupted or cannot be read
    """
    try:
        docx = DocxDocument(io.BytesIO(file_bytes))

        text_parts = []

        for paragraph in docx.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in docx.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)

        return "\n".join(text_parts)
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from DOCX: {e}") from e


def extract_from_csv(file_bytes: bytes) -> str:
    """Extract text from CSV file.

    Args:
        file_bytes: Raw CSV file bytes

    Returns:
        Formatted text with rows as lines

    Raises:
        ExtractionError: If CSV cannot be parsed
    """
    try:
        content = file_bytes.decode("utf-8", errors="replace")
        df = pd.read_csv(io.StringIO(content))

        text_parts = []
        for _, row in df.iterrows():
            row_text = ", ".join(str(v).strip() for v in row.values)
            text_parts.append(row_text)

        return "\n".join(text_parts)
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from CSV: {e}") from e


def extract_from_xlsx(file_bytes: bytes) -> str:
    """Extract text from XLSX file.

    Args:
        file_bytes: Raw XLSX file bytes

    Returns:
        Formatted text with sheet delimiters

    Raises:
        ExtractionError: If XLSX cannot be parsed
    """
    try:
        excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
        text_parts = []

        for sheet_name in excel_file.sheet_names:
            text_parts.append(f"=== {sheet_name} ===")

            df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name)

            for _, row in df.iterrows():
                row_text = ", ".join(str(v).strip() for v in row.values)
                if row_text.strip():
                    text_parts.append(row_text)

            text_parts.append("")

        return "\n".join(text_parts)
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from XLSX: {e}") from e


def extract_from_txt(file_bytes: bytes) -> str:
    """Extract text from TXT file.

    Args:
        file_bytes: Raw TXT file bytes

    Returns:
        Decoded text

    Raises:
        ExtractionError: If encoding fails
    """
    try:
        return file_bytes.decode("utf-8", errors="replace")
    except Exception as e:
        raise ExtractionError(f"Failed to extract text from TXT: {e}") from e


__all__ = [
    "ExtractionError",
    "extract_from_csv",
    "extract_from_docx",
    "extract_from_pdf",
    "extract_from_txt",
    "extract_from_xlsx",
]
